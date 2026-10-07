"""Savollar Excel formati: eksport, import va shablon bir tilda gaplashadi.

Nima uchun bu testlar bor. Ilgari eksport sarlavhalarni o'zicha yozardi
(«№ | Savol | A variant | …»), import esa ustunlarni o'rni bo'yicha o'qirdi.
Eksport qilingan faylni qaytadan yuklasa, savol matni o'rniga qator raqami
tushardi, variantlar bittaga surilardi va barcha to'g'ri javoblar «a» ga
aylanardi. Savollar yaratilaverardi — xato hech qayerda ko'rinmasdi.
"""

import io

import pytest
from openpyxl import Workbook, load_workbook

from app.modules.quiz.question.excel_format import (
    TEMPLATE_HEADERS,
    normalize_header,
    parse_correct_option,
    read_question_sheet,
    resolve_columns,
)

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _xlsx(rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


async def _upload(auth_client, subject_id: int, content: bytes):
    return await auth_client.post(
        f"/question/upload_excel?subject_id={subject_id}",
        files={"file": ("savollar.xlsx", content, XLSX_MIME)},
    )


@pytest.mark.asyncio
async def test_exported_file_can_be_imported_back(auth_client, test_subject):
    """Asosiy regressiya: eksport → import savollarni buzmaydi."""
    users_resp = await auth_client.get("/user/")
    user_id = users_resp.json()["users"][0]["id"]

    created = await auth_client.post(
        "/question/",
        json={
            "subject_id": test_subject.id,
            "user_id": user_id,
            "text": "2 + 2 nechaga teng?",
            "option_a": "3",
            "option_b": "4",
            "option_c": "5",
            "option_d": "6",
            "correct_option": "b",
        },
    )
    assert created.status_code == 201, created.json()

    exported = await auth_client.get("/question/download_excel", params={"subject_id": test_subject.id})
    assert exported.status_code == 200

    response = await _upload(auth_client, test_subject.id, exported.content)
    assert response.status_code == 201, response.json()

    questions = response.json()["questions"]
    assert len(questions) == 1
    imported = questions[0]
    # Ilgari bu yerda text «1» (qator raqami), option_a esa savol matni edi.
    assert imported["text"] == "2 + 2 nechaga teng?"

    # Endi to'g'ri javob ustuni yo'q: eksport javobni «A variant» ga olib
    # chiqadi, import esa A ni to'g'ri deb oladi. Shuning uchun HARF emas,
    # javob MATNI tekshiriladi — ma'no aynan shunda saqlanadi.
    assert imported["correct_option"] == "a"
    assert imported["option_a"] == "4", "to'g'ri javob («4») A ustunida qolishi kerak"
    # Qolgan variantlar yo'qolmaydi, faqat tartibi siljiydi.
    assert sorted(
        [imported["option_a"], imported["option_b"], imported["option_c"], imported["option_d"]]
    ) == ["3", "4", "5", "6"]


@pytest.mark.asyncio
async def test_legacy_positional_file_still_imports(auth_client, test_subject):
    """Sarlavhalar tanilmasa — eski, o'rni bo'yicha o'qish ishlaydi."""
    content = _xlsx(
        [
            ["Ustun1", "Ustun2", "Ustun3", "Ustun4", "Ustun5", "correct_option"],
            ["Poytaxt qaysi?", "Samarqand", "Toshkent", "Buxoro", "Xiva", "b"],
        ]
    )

    response = await _upload(auth_client, test_subject.id, content)
    assert response.status_code == 201, response.json()

    imported = response.json()["questions"][0]
    assert imported["text"] == "Poytaxt qaysi?"
    assert imported["option_a"] == "Samarqand"
    assert imported["correct_option"] == "b"


@pytest.mark.asyncio
async def test_headers_ignore_case_and_apostrophe_shape(auth_client, test_subject):
    """«TO'G'RI JAVOB» ham, «to'g'ri javob» ham bir xil tushuniladi.

    Excel apostrofni avtomatik «'» ga almashtiradi, foydalanuvchi esa
    qaysi belgi turganini ko'rmaydi.
    """
    content = _xlsx(
        [
            ["SAVOL", "a variant", "B Variant", "c variant", "D VARIANT", "TO’G’RI JAVOB"],
            ["Eng katta sayyora?", "Mars", "Yupiter", "Venera", "Saturn", "B"],
        ]
    )

    response = await _upload(auth_client, test_subject.id, content)
    assert response.status_code == 201, response.json()

    imported = response.json()["questions"][0]
    assert imported["text"] == "Eng katta sayyora?"
    assert imported["option_b"] == "Yupiter"
    assert imported["correct_option"] == "b"


@pytest.mark.asyncio
async def test_blank_rows_do_not_create_questions(auth_client, test_subject):
    """Savollar orasidagi bo'sh qatorlar savolga aylanmaydi.

    Bo'sh qatorlar ataylab oxirida emas, o'rtasida: oxirgilarini Excel'ning
    o'zi kesib tashlaydi va test hech narsani tekshirmagan bo'lardi.
    Ilgari bunday qator `text=""` bilan savol yaratardi.
    """
    content = _xlsx(
        [
            TEMPLATE_HEADERS,
            ["Birinchi savol", "A", "B", "C", "D", "a"],
            [None, None, None, None, None, None],
            ["", "", "", "", "", ""],
            ["Oxirgi savol", "A", "B", "C", "D", "b"],
        ]
    )

    response = await _upload(auth_client, test_subject.id, content)
    assert response.status_code == 201, response.json()

    questions = response.json()["questions"]
    assert [q["text"] for q in questions] == ["Birinchi savol", "Oxirgi savol"]


@pytest.mark.asyncio
async def test_new_template_marks_first_option_as_correct(auth_client, test_subject):
    """Yangi shablonda ustun yo'q: A — to'g'ri javob, ogohlantirishsiz.

    Ilgari bu holat «xato» deb ogohlantirilardi. Endi bu QOIDA, shuning
    uchun ogohlantirish ham yo'q — aks holda har bir qator uchun keraksiz
    ogohlantirish chiqib, haqiqiy muammolar ular orasida ko'rinmasdi.
    """
    content = _xlsx(
        [
            list(TEMPLATE_HEADERS),
            ["Javob A da", "To'g'ri", "Noto'g'ri 1", "Noto'g'ri 2", "Noto'g'ri 3"],
        ]
    )

    response = await _upload(auth_client, test_subject.id, content)
    assert response.status_code == 201, response.json()

    body = response.json()
    assert body["questions"][0]["correct_option"] == "a"
    assert body["questions"][0]["option_a"] == "To'g'ri"
    assert not body["warnings"], "qoida bo'yicha ishlagan faylga ogohlantirish kerak emas"


@pytest.mark.asyncio
async def test_legacy_file_with_correct_column_is_respected(auth_client, test_subject):
    """Eski fayldagi «To'g'ri javob» ustuni e'tiborsiz qolmaydi.

    O'qituvchilarda ilgari yuklab olingan fayllar bor va ularda haqiqiy
    «b»/«c»/«d» turadi. Agar ustun shunchaki tashlab yuborilsa, o'z
    bankini qayta yuklagan o'qituvchining BARCHA javoblari «a» bo'lib
    qolardi — bu jimgina ma'lumot buzilishi.
    """
    content = _xlsx(
        [
            ["Savol", "A variant", "B variant", "C variant", "D variant", "To'g'ri javob"],
            ["Poytaxt qaysi?", "Samarqand", "Toshkent", "Buxoro", "Xiva", "B"],
        ]
    )

    response = await _upload(auth_client, test_subject.id, content)
    assert response.status_code == 201, response.json()
    assert response.json()["questions"][0]["correct_option"] == "b"


@pytest.mark.asyncio
async def test_cyrillic_letter_in_legacy_file_is_understood(auth_client, test_subject):
    """Kirillcha «В» lotinchadan farq qilmaydi, lekin boshqa belgi.

    Ilgari bunday qiymat tanilmay, javob jimgina «a» bo'lardi.
    """
    content = _xlsx(
        [
            ["Savol", "A variant", "B variant", "C variant", "D variant", "To'g'ri javob"],
            ["Kirillcha javob", "Bir", "Ikki", "Uch", "To'rt", "В"],
        ]
    )

    response = await _upload(auth_client, test_subject.id, content)
    assert response.status_code == 201, response.json()
    assert response.json()["questions"][0]["correct_option"] == "b"


@pytest.mark.asyncio
async def test_template_has_headers_and_empty_first_sheet(auth_client):
    response = await auth_client.get("/question/excel_template")
    assert response.status_code == 200

    wb = load_workbook(io.BytesIO(response.content))
    sheet = wb["Savollar"]
    headers = [c.value for c in sheet[1]]
    assert len(headers) == 5, "«To'g'ri javob» ustuni shablondan olib tashlangan"
    # A ustuni sarlavhasida qoida yozilgan: o'qituvchi «Namuna» varag'ini
    # ochmasligi mumkin.
    assert headers[1].startswith(TEMPLATE_HEADERS[1])
    assert "to'g'ri javob" in headers[1].lower()
    assert headers[0] == TEMPLATE_HEADERS[0]
    # Birinchi varaqda ma'lumot bo'lmasligi shart: import faqat shuni
    # o'qiydi, namuna qatori bazaga tushib qolmasligi kerak.
    assert sheet.max_row == 1
    assert "Namuna" in wb.sheetnames
    assert wb["Namuna"].max_row > 1


@pytest.mark.asyncio
async def test_template_is_importable(auth_client, test_subject):
    """Shablonni to'ldirib yuklash ishlaydi — sarlavhalari tanildi."""
    template = await auth_client.get("/question/excel_template")
    wb = load_workbook(io.BytesIO(template.content))
    sheet = wb["Savollar"]
    sheet.append(["Shablondan savol", "To'g'ri", "Xato 1", "Xato 2", "Xato 3"])
    buffer = io.BytesIO()
    wb.save(buffer)

    response = await _upload(auth_client, test_subject.id, buffer.getvalue())
    assert response.status_code == 201, response.json()

    imported = response.json()["questions"][0]
    assert imported["text"] == "Shablondan savol"
    assert imported["correct_option"] == "a"
    assert imported["option_a"] == "To'g'ri"


def test_export_headers_are_all_recognized():
    """Eksport sarlavhalarini parser taniydi — ikkisi yana ajralmasin.

    Bu sof birlik testi: `download_questions_excel` dagi ro'yxat qo'lda
    takrorlangan, chunki u funksiya ichida turadi. Agar u o'zgarsa va
    alias qo'shilmasa, shu test yiqiladi — eksport → import yana jimgina
    buzilishidan oldin.
    """
    export_headers = [
        "№",
        "Savol",
        "A variant (to'g'ri javob)",
        "B variant",
        "C variant",
        "D variant",
        "Fan",
        "Foydalanuvchi",
    ]
    mapping = resolve_columns(export_headers)

    assert mapping is not None, "eksport sarlavhalari tanilmadi"
    assert mapping["text"] == 1
    # Qavs ichidagi izoh sarlavhani buzmasligi kerak: aks holda fayl
    # ustun O'RNI bo'yicha o'qilib, birinchi ustundagi «№» savol matni
    # bo'lib tushardi.
    assert mapping["option_a"] == 2
    assert mapping["option_d"] == 5


def test_unknown_headers_fall_back_to_positional():
    assert resolve_columns(["a", "b", "c", "d", "e"]) is None


def test_normalize_header_folds_apostrophes_and_spaces():
    assert normalize_header("  TO’G‘RI   JAVOB ") == "to'g'ri javob"


def test_parse_correct_option_understands_common_mistakes():
    """Eski fayllardagi yozuv shakllari — bitta joyda tekshiriladi."""
    assert parse_correct_option("B") == "b"
    assert parse_correct_option("  c  ") == "c"
    assert parse_correct_option("D)") == "d"
    assert parse_correct_option("2") == "b"
    # Kirillcha harflar lotinchadan ko'z bilan farq qilmaydi.
    assert parse_correct_option("В") == "b"
    assert parse_correct_option("С") == "c"
    assert parse_correct_option("B variant") == "b"
    assert parse_correct_option("variant b") == "b"
    # Tanib bo'lmaydigan qiymat — `None`, chaqiruvchi «a» qo'yadi.
    assert parse_correct_option("") is None
    assert parse_correct_option(None) is None
    # Matnning birinchi harfi bo'yicha taxmin qilinmaydi.
    assert parse_correct_option("bilmadim") is None
    # Ikki xil harf ko'rsatilgan — taxmin qilmaymiz.
    assert parse_correct_option("a yoki b") is None



# ─── Sarlavha qatori qayerda ─────────────────────────────────────────────────
#
# Ilgari `pd.read_excel` birinchi qatorni HAR DOIM sarlavha deb olardi:
# sarlavhasiz faylda birinchi savol jimgina yo'qolardi, sarlavha ustida nom
# yoki bo'sh qator bo'lsa — sarlavhaning o'zi savol bo'lib bankka tushardi.

_H = ["Savol", "A variant", "B variant", "C variant", "D variant"]
_Q1 = ["2+2?", "4", "3", "5", "6"]
_Q2 = ["3+3?", "6", "5", "7", "8"]


def _texts(sheet, position: int = 0) -> list[str]:
    index = sheet.mapping["text"] if sheet.mapping else position
    return [str(cells[index]) for _, cells in sheet.rows if any(v is not None for v in cells)]


@pytest.mark.parametrize(
    ("name", "rows", "expected_lines"),
    [
        ("sarlavha bilan", [_H, _Q1, _Q2], [2, 3]),
        ("sarlavhasiz", [_Q1, _Q2], [1, 2]),
        ("sarlavha ustida nom", [["Matematika savollari"], _H, _Q1, _Q2], [3, 4]),
        ("birinchi qator bo'sh", [[], _H, _Q1, _Q2], [3, 4]),
    ],
)
def test_header_row_is_found_wherever_it_is(name, rows, expected_lines):
    sheet = read_question_sheet(_xlsx(rows))

    assert _texts(sheet) == ["2+2?", "3+3?"], name
    assert [line for line, cells in sheet.rows if any(v is not None for v in cells)] == expected_lines, name


def test_unknown_but_header_like_first_row_is_skipped_with_note():
    sheet = read_question_sheet(_xlsx([["Savollar", "1-variant", "2-variant", "3-variant", "4-variant"], _Q1, _Q2]))

    assert sheet.mapping is None
    assert _texts(sheet) == ["2+2?", "3+3?"]
    assert sheet.notes == ["1-qator sarlavha deb hisoblandi va savol sifatida yuklanmadi"]


def test_question_mentioning_variant_is_not_mistaken_for_header():
    """Bitta katakda «variant» so'zi — bu savol, sarlavha emas."""
    question = ["Qaysi variant to'g'ri?", "Birinchi", "Ikkinchi", "Uchinchi", "To'rtinchi"]
    sheet = read_question_sheet(_xlsx([question, _Q1]))

    assert _texts(sheet) == ["Qaysi variant to'g'ri?", "2+2?"]
    assert sheet.notes == []


@pytest.mark.asyncio
async def test_upload_without_header_keeps_first_question(auth_client, test_subject):
    response = await _upload(auth_client, test_subject.id, _xlsx([_Q1, _Q2]))

    assert response.status_code == 201, response.json()
    assert [q["text"] for q in response.json()["questions"]] == ["2+2?", "3+3?"]


@pytest.mark.asyncio
async def test_upload_with_title_row_does_not_import_header(auth_client, test_subject):
    content = _xlsx([["Matematika savollari"], [], _H, _Q1, _Q2])

    response = await _upload(auth_client, test_subject.id, content)

    assert response.status_code == 201, response.json()
    assert [q["text"] for q in response.json()["questions"]] == ["2+2?", "3+3?"]


@pytest.mark.asyncio
async def test_general_test_upload_finds_header_and_real_line_numbers(auth_client):
    """Elementar test ham xuddi shu o'quvchidan foydalanadi; ogohlantirishdagi
    qator raqami — Excel'dagi haqiqiy raqam."""
    subject = await auth_client.post("/general-test/subject", json={"name": "Excel sarlavha fani"})
    assert subject.status_code == 201, subject.text
    content = _xlsx([["Sarlavha ustidagi nom"], _H, _Q1, ["Chala savol", "4", "", "5", "6"], _Q2])

    response = await auth_client.post(
        f"/general-test/subject/{subject.json()['id']}/upload_excel",
        files={"file": ("savollar.xlsx", content, XLSX_MIME)},
    )

    assert response.status_code in (200, 201), response.text
    body = response.json()
    assert body["created"] == 2
    assert any(w.startswith("4-qator") for w in body["warnings"]), body["warnings"]
